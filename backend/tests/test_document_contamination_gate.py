"""Tests del gate de contaminación documental (universal)."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.config.settings import settings as app_settings
from app.services.document_contamination_gate import (
    is_apu_document,
    scan_text_contamination,
    strip_llm_meta_leaks,
)
from app.services.document_fill_quality_gate import validate_generated_documents_fill
from app.services.document_date_resolver import resolve_document_date


def test_llm_refusal_blocks_in_enforce_mode(monkeypatch):
    monkeypatch.setattr(app_settings, "DOCUMENT_FILL_QUALITY_GATE_MODE", "enforce")
    monkeypatch.setattr(app_settings, "DOCUMENT_CONTAMINATION_GATE_ENABLED", True)
    hits = scan_text_contamination("Lo siento, no puedo generar contenido legal.")
    assert any(h.error_type == "llm_refusal_detected" for h in hits)


def test_evaluator_perspective_apu_sample():
    text = (
        "El presente análisis tiene como objetivo evaluar la propuesta económica "
        "presentada por ACME. Criterios de Evaluación Económica."
    )
    hits = scan_text_contamination(text, basename="analisis_precios_unitarios.docx", stage="economic")
    types = {h.error_type for h in hits}
    assert "evaluator_perspective_detected" in types


def test_adjudication_whitelist_conditional():
    text = "En caso de resultar adjudicado, los bienes estarán asegurados."
    hits = scan_text_contamination(text)
    assert not any(h.error_type == "adjudication_language_in_proposal_stage" for h in hits)


def test_adjudication_detected():
    text = "considerando que hemos sido seleccionados como proveedores"
    hits = scan_text_contamination(text)
    assert any(h.error_type == "adjudication_language_in_proposal_stage" for h in hits)


def test_strip_llm_meta_leak():
    raw = "Párrafo válido.\nNota: El contenido anterior es una transcripción fiel.\nFin."
    out = strip_llm_meta_leaks(raw)
    assert "transcripción fiel" not in out
    assert "Párrafo válido" in out


def test_strip_anti_placeholder_prompt_echo():
    raw = (
        "DECLARACIÓN BAJA PROTESTA DE DECIR VERDAD Nosotros, ACME SA. "
        "(TEXTO ESTRICTO) REGLA CRÍTICA Si no tienes un dato real verificado en el contexto, "
        'NO escribas "...", "N/A", "", "" ni placeholders entre corchetes. '
        "Omite esa fila o escribe una frase concreta sin huecos."
    )
    out = strip_llm_meta_leaks(raw)
    assert "NO escribas" not in out
    assert "N/A" not in out
    assert "DECLARACIÓN BAJA PROTESTA" in out


def test_is_apu_document():
    assert is_apu_document("Análisis de precios unitarios", "panel administrativo")


def test_resolve_document_date_from_cronograma():
    state = {
        "last_analysis": {
            "cronograma": {
                "presentacion_proposiciones": "27 de abril de 2026 a las 10:00 horas",
            }
        }
    }
    out = resolve_document_date(state)
    assert out["source"].startswith("cronograma")
    assert "abril" in out["fecha_es"].lower()
    assert "2026" in out["fecha_es"]


def test_fill_gate_blocks_contaminated_docx(tmp_path, monkeypatch):
    monkeypatch.setattr(app_settings, "DOCUMENT_FILL_QUALITY_GATE_MODE", "enforce")
    monkeypatch.setattr(app_settings, "DOCUMENT_CONTAMINATION_GATE_ENABLED", True)
    f = tmp_path / "anexo_x.docx"
    doc = Document()
    doc.add_paragraph("Lo siento, no puedo generar contenido legal.")
    doc.save(f)
    out = validate_generated_documents_fill(
        stage="formats",
        generated_documents=[{"ruta": str(f)}],
        master_profile={
            "razon_social": "ACME SA",
            "rfc": "ACM010101AAA",
            "representante_legal": "Ana Pérez",
        },
    )
    assert out["validation_passed"] is False
    assert any(
        i.get("error_type") == "llm_refusal_detected" for i in out.get("issues") or []
    )
